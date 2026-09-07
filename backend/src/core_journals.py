# -*- coding: utf-8 -*-
"""核心期刊清单加载与匹配。

数据源(DB 主存,启动期冷填充):
  表 core_journals (id, list_source, name, name_key, created_at)
  - list_source ∈ {pku_core, cscd_core, cscd_ext, cssci_core}

匹配键(name_key):去装饰符号、全/半角统一、小写,严格相等。

启动行为:
  - init_db() 会建表;
  - 第一次启动若表为空,从 xlsx(若存在)冷导入;
  - 已导入过则直接走内存缓存,不读 xlsx。

后端运行期可用 /api/core-journals/reload 强制重导入。
"""
from __future__ import annotations

import logging
import os
import re
from functools import lru_cache
from pathlib import Path

import openpyxl

log = logging.getLogger(__name__)

# 数据库 Session 延迟导入,避免循环
def _get_session():
    from db.session import SessionLocal
    from db.models import CoreJournalModel
    return SessionLocal, CoreJournalModel


# === xlsx 来源路径(只在冷启动时使用一次) ===

BASE_DIR = Path(__file__).resolve().parents[2]
PKU_XLSX = Path(os.environ.get(
    "PKU_CORE_XLSX",
    BASE_DIR / "北京大学《中文核心期刊要目总览》（2023年版）.xlsx",
))
CSCD_XLSX = Path(os.environ.get(
    "CSCD_XLSX",
    BASE_DIR / "中国科学引文数据库来源期刊列表(2025-2026).xlsx",
))
CSSCI_XLSX = Path(os.environ.get(
    "CSSCI_XLSX",
    BASE_DIR / "CSSCI 来源期刊列表.xlsx",
))


# === 名称归一化 ===

_DECORATE_RE = re.compile(r"[\s\u3000《》<>()()【】\[\]「」『』\-—–·,。.,、/／:：]+")


def _normalize(name: str | None) -> str:
    if not name:
        return ""
    s = name.strip()
    s = s.translate(str.maketrans({
        "（": "(", "）": ")", "，": ",", "。": ".", "：": ":",
        "；": ";", "？": "?", "！": "!", "／": "/", "—": "-", "–": "-",
        "＋": "+", "＝": "=", "＠": "@", "＆": "&",
        "“": '"', "”": '"', "‘": "'", "’": "'",
    }))
    s = s.lower()
    s = _DECORATE_RE.sub("", s)
    return s


def clean_journal_name(raw: str | None) -> str:
    """清洗抓取侧粘在刊名上的年卷期/收录来源尾巴。

    知网详情页 source 字段实测污染形态:
      '中国安全生产科学技术 . 2026\\n ,22\\n (07) 查看该刊数据库收录来源'
      '中国城市规划知识仓库 . 查看该刊数据库收录来源'
    规则: ①截断「查看该刊数据库收录来源」及之后全部内容;
          ②按 " . "(空格点空格,知网拼接年卷期的分隔符)取首段;
          ③strip 收尾。
    """
    if not raw:
        return ""
    s = raw
    idx = s.find("查看该刊数据库收录来源")
    if idx >= 0:
        s = s[:idx]
    if " . " in s:
        s = s.split(" . ", 1)[0]
    return s.strip()


# === xlsx 读取(供冷启动时一次性灌入 DB) ===

def _load_pku_core(path: Path) -> list[str]:
    if not path.is_file():
        return []
    out: list[str] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 2:
            continue
        num, name = row[0], row[1]
        if not isinstance(num, (int, float)):
            continue
        if not name or not isinstance(name, str):
            continue
        out.append(name.strip())
    wb.close()
    return out


def _load_cscd(path: Path) -> tuple[list[str], list[str]]:
    core: list[str] = []
    ext: list[str] = []
    if not path.is_file():
        return core, ext
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 4:
            continue
        num, name, _issn, tag = row[0], row[1], row[2], row[3]
        if not isinstance(num, (int, float)):
            continue
        if not name or not isinstance(name, str):
            continue
        if isinstance(tag, str) and "扩展" in tag:
            ext.append(name.strip())
        else:
            core.append(name.strip())
    wb.close()
    return core, ext


def _load_cssci(path: Path) -> list[str]:
    if not path.is_file():
        return []
    out: list[str] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        for cell in row:
            if isinstance(cell, str):
                stripped = cell.strip()
                if stripped:
                    out.append(stripped)
                    break
    wb.close()
    return out


# === DB 灌入 ===

def _seed_from_xlsx_if_empty() -> None:
    """第一次启动:把 xlsx 灌进 core_journals 表(只在表为空时执行)。

    v7.1:先查表再读 xlsx。xlsx 已被删除(数据已于首次启动导入落库),
    表非空时不再尝试读 xlsx,避免误报「xlsx 全部缺失」警告。
    """
    SessionLocal, CoreJournalModel = _get_session()
    with SessionLocal() as s:
        try:
            existing = s.query(CoreJournalModel).limit(1).count()
        except Exception:
            existing = 0
        if existing > 0:
            # 数据已落库(xlsx 已永久删除,表即唯一数据源),静默跳过
            return

    # === 提前:同一进程内去重(同一 key 不重复入 list_source)===
    seen_per_src: dict[str, set[str]] = {
        "pku_core": set(), "cscd_core": set(), "cscd_ext": set(), "cssci_core": set(),
    }
    rows: list[CoreJournalModel] = []

    def _add(src: str, name: str) -> None:
        key = _normalize(name)
        if not key or len(key) < 2:
            return
        if key in seen_per_src[src]:
            return
        seen_per_src[src].add(key)
        rows.append(CoreJournalModel(list_source=src, name=name, name_key=key))

    for name in _load_pku_core(PKU_XLSX):
        _add("pku_core", name)
    cscd_core_names, cscd_ext_names = _load_cscd(CSCD_XLSX)
    for name in cscd_core_names:
        _add("cscd_core", name)
    for name in cscd_ext_names:
        _add("cscd_ext", name)
    for name in _load_cssci(CSSCI_XLSX):
        _add("cssci_core", name)

    if not rows:
        log.warning("xlsx 缺失且 core_journals 表为空,核心期刊筛选不可用")
        return

    with SessionLocal() as s:
        try:
            existing = s.query(CoreJournalModel).limit(1).count()
        except Exception:
            existing = 0
        if existing > 0:
            return
        s.add_all(rows)
        try:
            s.commit()
            log.info("core_journals 冷导入完成: %d 条", len(rows))
        except Exception as exc:
            s.rollback()
            # 兜底:可能是跨进程残留 -> 用 IGNORE 重试
            try:
                from sqlalchemy import text
                with s.begin():
                    s.execute(
                        text(
                            "INSERT OR IGNORE INTO core_journals "
                            "(list_source, name, name_key, created_at) "
                            "VALUES (:src, :name, :key, CURRENT_TIMESTAMP)"
                        ),
                        [
                            {"src": r.list_source, "name": r.name, "key": r.name_key}
                            for r in rows
                        ],
                    )
                log.info("core_journals 冷导入(IGNORE)完成: %d 条", len(rows))
            except Exception as exc2:
                log.error("core_journals 冷导入失败(两条路都挂): %s | %s", exc, exc2)


def force_reseed_from_xlsx() -> int:
    """删表后重新从 xlsx 灌入(管理接口:reload 走这条)。"""
    SessionLocal, CoreJournalModel = _get_session()
    with SessionLocal() as s:
        s.query(CoreJournalModel).delete()
        s.commit()
    _seed_from_xlsx_if_empty()
    invalidate_cache()
    return stats()["all"]


# === 内存缓存(按 list_source 集合) ===

@lru_cache(maxsize=1)
def _load_all() -> dict[str, set[str]]:
    """从 DB 加载所有 list_source 的 name_key 集合。

    若表为空,尝试一次冷导入(通常 main.py 启动时已导入,这里再保一手)。
    """
    SessionLocal, CoreJournalModel = _get_session()
    out: dict[str, set[str]] = {
        "pku_core": set(),
        "cscd_core": set(),
        "cscd_ext": set(),
        "cssci_core": set(),
    }
    try:
        with SessionLocal() as s:
            for src in out.keys():
                rows = s.query(CoreJournalModel).filter(
                    CoreJournalModel.list_source == src
                ).all()
                for r in rows:
                    out[src].add(r.name_key)
    except Exception as exc:
        log.error("读 core_journals 失败: %s", exc)
        return out
    total = sum(len(v) for v in out.values())
    if total == 0:
        log.info("core_journals 空,尝试冷导入")
        try:
            _seed_from_xlsx_if_empty()
            with SessionLocal() as s:
                for src in out.keys():
                    rows = s.query(CoreJournalModel).filter(
                        CoreJournalModel.list_source == src
                    ).all()
                    for r in rows:
                        out[src].add(r.name_key)
        except Exception as exc:
            log.error("冷导入失败: %s", exc)
    log.info(
        "核心期刊加载: 北大核心=%d, CSCD 核心=%d, CSCD 扩展=%d, CSSCI=%d",
        len(out["pku_core"]), len(out["cscd_core"]),
        len(out["cscd_ext"]), len(out["cssci_core"]),
    )
    return out


def classify_journal(journal: str | None) -> dict:
    # v7.3:入口先清洗 —— 知网抓取的刊名带「 . 年卷期 查看该刊数据库收录来源」
    # 脏尾巴,不清洗则精确匹配永远打不中,「仅核心」筛选恒为 0 条。
    n = _normalize(clean_journal_name(journal))
    if not n:
        return {
            "pku_core": False, "cscd_core": False, "cscd_ext": False,
            "cssci_core": False, "in_core": False,
        }
    data = _load_all()
    return {
        "pku_core": n in data["pku_core"],
        "cscd_core": n in data["cscd_core"],
        "cscd_ext": n in data["cscd_ext"],
        "cssci_core": n in data["cssci_core"],
        "in_core": any(n in v for v in data.values()),
    }


def stats() -> dict:
    data = _load_all()
    return {
        "pku_core": len(data["pku_core"]),
        "cscd_core": len(data["cscd_core"]),
        "cscd_ext": len(data["cscd_ext"]),
        "cssci_core": len(data["cssci_core"]),
        "all": sum(len(v) for v in data.values()),
    }


def invalidate_cache() -> None:
    """强制重载(管理接口用,xlsx 替换或 reload 后调用)。"""
    _load_all.cache_clear()


def init_db() -> None:
    """应用启动时调用:首次从 xlsx 灌入 core_journals。"""
    _seed_from_xlsx_if_empty()
